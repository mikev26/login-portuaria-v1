import sys
import openpyxl
from io import BytesIO

from bitacora.services.practicaje_excel import crear_excel_practicaje

# Mock records
sample_registros = [
    {
        "nro": 1,
        "registro": "REG-001",
        "buque": "BUQUE TEST",
        "bandera": "ECUADOR",
        "ipb": "100",
        "ano": 2025,
        "ubicacion": "MANTA",
        "usuario": "ADMIN",
        "operadora": "OPERADORA TEST",
        "tipomaniobra": "ATRAQUE",
        "trimestre": 1,
        "fecha": "2025-01-15",
    }
]

excel_bytes = crear_excel_practicaje(sample_registros, sano="2025")
wb = openpyxl.load_workbook(BytesIO(excel_bytes))
ws = wb.active

print("Sheet Title:", ws.title)
print("A1 Value:", ws["A1"].value)
print("A1 Font Size:", ws["A1"].font.size)
print("Row 1 Height:", ws.row_dimensions[1].height)
print("Row 2 (Headers):", [ws.cell(row=2, column=col).value for col in range(1, 13)])
print("Row 2 Height:", ws.row_dimensions[2].height)
print("Row 3 (Sample Data):", [ws.cell(row=3, column=col).value for col in range(1, 13)])
print("Freeze Panes:", ws.freeze_panes)
print("AutoFilter Ref:", ws.auto_filter.ref)

excel_bytes_2026 = crear_excel_practicaje(sample_registros, sano="2026")
wb2 = openpyxl.load_workbook(BytesIO(excel_bytes_2026))
ws2 = wb2.active
print("A1 Value (2026):", ws2["A1"].value)
