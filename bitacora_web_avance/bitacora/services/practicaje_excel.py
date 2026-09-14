"""Generación del archivo Excel del módulo de practicaje."""

from __future__ import annotations

import io
from typing import Any


COLUMNAS_PRACTICAJE = (
    ("Nro", "nro"),
    ("Registro", "registro"),
    ("Buque", "buque"),
    ("Bandera", "bandera"),
    ("IPB", "ipb"),
    ("Año", "ano"),
    ("Ubicación", "ubicacion"),
    ("Usuario", "usuario"),
    ("Operadora", "operadora"),
    ("TipoManiobra", "tipomaniobra"),
    ("Trimestre", "trimestre"),
    ("Fecha", "fecha"),
)


def crear_excel_practicaje(
    registros: list[dict[str, Any]],
    sano: str = "",
) -> bytes:
    """Construye el Excel con el encabezado dinámico y filas formateadas."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "datos_practicaje"

    # Determinar el año seleccionado
    anio_str = str(sano).strip()
    if not anio_str and registros:
        for r in registros:
            if r.get("ano"):
                anio_str = str(r["ano"]).strip()
                break
    if not anio_str:
        anio_str = "2025"

    # Fila 1: Encabezado principal del servicio de practicaje
    titulo_encabezado = (
        f"SERVICIO DE PRACTICAJE  AÑO {anio_str} - TPyC y TERMINALES PRIVADOS EXCEPTO TPM."
    )
    total_columnas = len(COLUMNAS_PRACTICAJE)
    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=total_columnas,
    )
    fill_azul = PatternFill("solid", fgColor="1F4E79")
    for col in range(1, total_columnas + 1):
        c = worksheet.cell(row=1, column=col)
        c.fill = fill_azul
        c.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.cell(row=1, column=1, value=titulo_encabezado)
    worksheet.row_dimensions[1].height = 34

    # Fila 2: Encabezados de columnas
    encabezados = [titulo for titulo, _ in COLUMNAS_PRACTICAJE]
    worksheet.append(encabezados)

    worksheet.row_dimensions[2].height = 26
    for cell in worksheet[2]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Filas 3+: Datos
    for registro in registros:
        worksheet.append([registro.get(clave, "") for _, clave in COLUMNAS_PRACTICAJE])

    worksheet.freeze_panes = "A3"
    col_max_letra = get_column_letter(total_columnas)
    total_filas = len(registros) + 2
    worksheet.auto_filter.ref = f"A2:{col_max_letra}{total_filas}"

    # Ajuste de ancho de columnas (omitimos la fila 1 para no distorsionar por el título largo)
    for col_idx in range(1, total_columnas + 1):
        col_letter = get_column_letter(col_idx)
        longest = max(
            len(str(worksheet.cell(row=r, column=col_idx).value or ""))
            for r in range(2, len(registros) + 3)
        )
        worksheet.column_dimensions[col_letter].width = min(
            max(longest + 4, 14), 38
        )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
