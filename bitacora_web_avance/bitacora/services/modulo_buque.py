"""Servicios para el módulo de buques."""

from __future__ import annotations

from typing import Any

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.styles import Font

from .db_connection import execute_query, execute_procedure


def obtener_buques() -> list[dict[str, Any]]:
    """Obtiene el listado general de buques."""

    return execute_query(
        "EXEC dbo.SPJ_InfoBuques"
    )


def obtener_registros(id_buque: Any) -> list[dict[str, Any]]:
    """Obtiene los registros asociados a un buque."""

    return execute_query(
        "EXEC dbo.SPJ_consulta_registros ?",
        (id_buque,),
    )


def obtener_operadores_movimiento(
    solicitud_param: Any,
) -> list[dict[str, Any]]:
    """
    Obtiene los operadores asociados a una solicitud.

    Si la solicitud viene en formato SCANUAL, primero
    obtiene el idsolicitud correspondiente.
    """

    solicitud_param = str(solicitud_param or "").strip()

    if "-" in solicitud_param:
        rows = execute_query(
            """
            SELECT idsolicitud
            FROM dbo.dim_mov_solicitud
            WHERE scanual = ?
            """,
            (solicitud_param,),
        )

        solicitud_val = (
            rows[0].get("idsolicitud")
            if rows
            else 0
        )
    else:
        try:
            solicitud_val = int(solicitud_param) if solicitud_param else 0
        except (TypeError, ValueError):
            solicitud_val = 0

    return execute_query(
        "EXEC dbo.SPJ_consulta_mov_operadores ?",
        (solicitud_val,),
    )


def obtener_operadores_listados() -> list[dict[str, Any]]:
    """Obtiene el listado general de operadores."""

    return execute_query(
        """
        DECLARE @res INT;

        EXEC dbo.SP_Operadores_Listados
            @sresult = @res OUTPUT;
        """
    )


def _get_col(
    row_dict: dict[str, Any],
    *nombres_posibles: str,
) -> Any:
    """
    Obtiene un valor del diccionario ignorando mayúsculas/minúsculas
    en el nombre de la columna.
    """

    for nombre in nombres_posibles:
        for key, value in row_dict.items():
            if key.lower() == nombre.lower():
                return value if value is not None else ""

    return ""


def _obtener_buques_puerto(
    ubicacion_id: int,
) -> list[dict[str, Any]]:
    """Obtiene los buques según su ubicación en puerto."""

    return execute_query(
        "EXEC dbo.SPJ_BuquesPuerto ?",
        (ubicacion_id,),
    )


def exportar_buques(
    ubicacion_param: str,
    tipo_info: str,
) -> tuple[bytes, str]:
    """
    Genera el Excel del reporte de buques.

    Retorna:
        (contenido_archivo, nombre_archivo)
    """

    ubicacion_param = str(
        ubicacion_param or ""
    ).strip().lower()

    tipo_info = str(
        tipo_info or "listado"
    ).strip().lower()

    nombre_plantilla = (
        "plantilla_listado.xlsx"
        if tipo_info == "listado"
        else "plantilla_detalle.xlsx"
    )

    ruta_plantilla = (
        settings.BASE_DIR
        / "plantilla"
        / nombre_plantilla
    )

    if not ruta_plantilla.exists():
        raise FileNotFoundError(
            f"No se encontró la plantilla '{nombre_plantilla}'."
        )

    wb = load_workbook(ruta_plantilla)
    ws = wb.active

    fecha_actual = __import__("datetime").datetime.now().strftime(
        "%d/%m/%Y %H:%M"
    )

    ws["B6"] = fecha_actual
    ws["B6"].font = Font(
        color="000000",
        name="Calibri",
        size=11,
        bold=False,
    )

    if ubicacion_param == "muelle":
        ubicacion_id = 0
    elif ubicacion_param == "fondeo":
        ubicacion_id = 1
    else:
        ubicacion_id = 2

    resultados = _obtener_buques_puerto(
        ubicacion_id
    )

    fuente_negra = Font(
        color="000000",
        name="Calibri",
        size=11,
        bold=False,
    )

    fila_actual = 9

    for row in resultados:
        ub_bd = str(
            _get_col(
                row,
                "Ubicacion",
                "ubicacion",
            )
        ).strip().lower()

        if ubicacion_param == "muelle":
            condicion = (
                "muelle" in ub_bd
                or "abarloado" in ub_bd
                or "marginal" in ub_bd
            )
        else:
            condicion = (
                "fondeo" in ub_bd
                or "fondeadero" in ub_bd
            )

        if condicion or not ub_bd:

            if tipo_info == "listado":
                valores = [
                    _get_col(row, "Solicitud"),
                    _get_col(row, "Registro"),
                    _get_col(
                        row,
                        "buque",
                        "nombre_buque",
                    ),
                    _get_col(row, "Matricula"),
                    _get_col(row, "bandera"),
                    _get_col(row, "Eslora"),
                    _get_col(row, "TRB"),
                    _get_col(row, "TRN"),
                    _get_col(
                        row,
                        "Ubicacion",
                    ),
                ]

            else:
                valores = [
                    _get_col(row, "Solicitud"),
                    _get_col(row, "Registro"),
                    _get_col(
                        row,
                        "buque",
                        "nombre_buque",
                    ),
                    _get_col(row, "Matricula"),
                    _get_col(row, "bandera"),
                    _get_col(row, "agencia"),
                    _get_col(row, "armador"),
                    _get_col(row, "TipoNave"),
                    _get_col(row, "Contrato"),
                    _get_col(row, "Eslora"),
                    _get_col(row, "TRB"),
                    _get_col(row, "TRN"),
                    _get_col(row, "Calado"),
                    _get_col(
                        row,
                        "Manga",
                        "MAnga",
                    ),
                    _get_col(row, "arribo"),
                    _get_col(
                        row,
                        "Ubicacion",
                    ),
                ]

            for col_idx, value in enumerate(
                valores,
                start=1,
            ):
                celda = ws.cell(
                    row=fila_actual,
                    column=col_idx,
                    value=value,
                )
                celda.font = fuente_negra

            fila_actual += 1

    import io

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nombre_archivo = (
        f"Reporte_Buques_"
        f"{ubicacion_param.capitalize()}_"
        f"{tipo_info.capitalize()}.xlsx"
    )

    return (
        output.getvalue(),
        nombre_archivo,
    )