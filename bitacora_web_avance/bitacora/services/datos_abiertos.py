from __future__ import annotations

from typing import Any

from .db_connection import (
    DatabaseConfigurationError,
    DatabaseContractError,
    _IDENTIFIER_RE,
    execute_procedure,
)


DATOS_ABIERTOS_FIELD_MAP = (
    ("REGISTRO", "Registro"),
    ("CODBUQUE", "CodBuque"),
    ("MATRÍCULA", "Matrícula"),
    ("BUQUE", "Buque"),
    ("TipoNave", "Tipo de Nave"),
    ("Arribo", "Arribo"),
    ("Zarpe", "Zarpe"),
    ("Bandera", "Bandera"),
    ("TRB", "TRB"),
    ("TRN", "TRN"),
    ("Agencia", "Agencia"),
    ("TotalDescarga", "Total Descarga"),
)


def _coerce_row_value(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in row and row[key] is not None:
            return row[key]

    normalized = {str(k).lower(): v for k, v in row.items()}
    for key in keys:
        normalized_key = str(key).lower()
        if normalized_key in normalized and normalized[normalized_key] is not None:
            return normalized[normalized_key]
    return None


def _map_datos_abiertos_row(row: dict[str, Any]) -> dict[str, Any]:
    mapped: dict[str, Any] = {}
    for source_key, target_key in DATOS_ABIERTOS_FIELD_MAP:
        value = _coerce_row_value(
            row,
            source_key,
            source_key.lower(),
            source_key.replace(" ", "_"),
            source_key.replace("-", "_"),
        )
        mapped[target_key] = value

        safe_key = target_key.replace(" ", "").replace("-", "")
        mapped[safe_key] = value

        if " " in target_key:
            mapped[target_key.replace(" ", "_")] = value

        if "Tipo de Nave" == target_key:
            mapped["TipoNave"] = value
        if "Total Descarga" == target_key:
            mapped["TotalDescarga"] = value

    return mapped


def obtener_reporte_datos_abiertos(
    anio: int,
    semestre: str | int,
) -> list[dict[str, Any]]:
    """Ejecución del SP dbo.SPJ_DatosAbiertosTPyC y mapeo de 12 columnas."""

    procedure_candidates = [
        "dbo.SPJ_DatosAbiertosTPyC",
        "dbo.SP_DatosAbiertosTPyC",
    ]

    semestre_valor = semestre
    if isinstance(semestre_valor, str):
        semestre_normalizado = str(semestre_valor or "").strip().lower()
        if semestre_normalizado in {"1er", "1"}:
            semestre_valor = 1
        elif semestre_normalizado in {"2do", "2"}:
            semestre_valor = 2
        else:
            raise DatabaseContractError(
                "El semestre debe ser '1er' o '2do' para datos abiertos."
            )

    semestre_numero = int(semestre_valor)
    if semestre_numero not in {1, 2}:
        raise DatabaseContractError(
            "El semestre debe ser '1er' o '2do' para datos abiertos."
        )

    procedure_params = {
        "dbo.SPJ_DatosAbiertosTPyC": (
            ("@sPeriodo", int(anio)),
            ("@sSemestre", semestre_numero),
        ),
        "dbo.SP_DatosAbiertosTPyC": (
            ("@sPeriodo", int(anio)),
            ("@sSemestre", semestre_numero),
        ),
    }

    last_error: Exception | None = None
    for procedure in procedure_candidates:
        if not _IDENTIFIER_RE.fullmatch(procedure):
            continue
        try:
            rows = execute_procedure(procedure, procedure_params[procedure])
            return [_map_datos_abiertos_row(row) for row in rows]
        except Exception as exc:  # pragma: no cover - compatibility fallback
            last_error = exc

    if last_error is not None:
        raise last_error

    raise DatabaseConfigurationError(
        "El procedimiento configurado para datos abiertos no tiene un formato válido."
    )


def generar_excel_datos_abiertos(
    registros: list[dict[str, Any]],
    template_path: str | Path,
    output_dir: str | Path | None = None,
    fecha_emision: Any = None,
) -> bytes:
    """Genera el archivo Excel de Datos Abiertos a partir de la plantilla institucional intacta.

    Conserva el encabezado, formato de celda, bordes, alineaciones y la fecha de emisión.
    No sobrescribe la plantilla original.
    """
    import copy
    from datetime import date, datetime
    import io
    import os
    from pathlib import Path
    import openpyxl

    template_str = os.fspath(template_path) if template_path else ""

    if not template_str or not os.path.exists(template_str):
        raise FileNotFoundError(
            f"No se encontró la plantilla Excel en la ruta configurada: {template_str}"
        )

    with open(template_str, "rb") as f:
        template_bytes = io.BytesIO(f.read())

    wb = openpyxl.load_workbook(template_bytes)
    ws = wb.active

    # Fecha de emisión en celda D6
    if isinstance(fecha_emision, (date, datetime)):
        fecha_str = fecha_emision.strftime("%d/%m/%Y")
    elif isinstance(fecha_emision, str) and fecha_emision.strip():
        fecha_str = fecha_emision.strip()
    else:
        fecha_str = date.today().strftime("%d/%m/%Y")

    ws.cell(6, 1).value = "Fecha de Emisión :"
    ws.cell(6, 4).value = f"Manta, {fecha_str}"

    start_row = 9
    template_cells = [ws.cell(9, col) for col in range(1, 13)]

    for idx, reg in enumerate(registros, start=start_row):
        row_values = [
            reg.get("Registro") if reg.get("Registro") is not None else reg.get("REGISTRO"),
            reg.get("CodBuque") if reg.get("CodBuque") is not None else reg.get("CODBUQUE"),
            reg.get("Matrícula") if reg.get("Matrícula") is not None else reg.get("Matricula") if reg.get("Matricula") is not None else reg.get("MATRÍCULA"),
            reg.get("Buque") if reg.get("Buque") is not None else reg.get("BUQUE"),
            reg.get("TipoNave") if reg.get("TipoNave") is not None else reg.get("Tipo de Nave"),
            reg.get("Arribo") if reg.get("Arribo") is not None else reg.get("ARRIBO"),
            reg.get("Zarpe") if reg.get("Zarpe") is not None else reg.get("ZARPE"),
            reg.get("Bandera") if reg.get("Bandera") is not None else reg.get("BANDERA"),
            reg.get("TRB") if reg.get("TRB") is not None else reg.get("trb"),
            reg.get("TRN") if reg.get("TRN") is not None else reg.get("trn"),
            reg.get("Agencia") if reg.get("Agencia") is not None else reg.get("AGENCIA"),
            reg.get("TotalDescarga") if reg.get("TotalDescarga") is not None else reg.get("Total Descarga"),
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=idx, column=col_idx)
            ref_cell = template_cells[col_idx - 1]

            cell.value = val
            if ref_cell.has_style:
                from openpyxl.styles import Font
                # Copy font but force color to black and ensure it's not bold
                original_font = ref_cell.font
                cell.font = Font(
                    name=original_font.name,
                    size=original_font.size,
                    bold=False,  # Asegurar que el texto normal no esté en negrita
                    italic=original_font.italic,
                    vertAlign=original_font.vertAlign,
                    underline=original_font.underline,
                    strike=original_font.strike,
                    color="FF000000"
                )
                cell.alignment = copy.copy(ref_cell.alignment)
                cell.border = copy.copy(ref_cell.border)
                cell.fill = copy.copy(ref_cell.fill)
                cell.number_format = ref_cell.number_format

    last_row = start_row + len(registros) - 1

    # Limpiar celdas excedentes en filas pre-formateadas si hay menos registros que la plantilla
    if ws.max_row > last_row:
        for r in range(last_row + 1, ws.max_row + 1):
            for col in range(1, 13):
                ws.cell(r, col).value = None

    output_buffer = io.BytesIO()
    wb.save(output_buffer)
    excel_bytes = output_buffer.getvalue()

    # Si se configuró un directorio de salida opcional, guardar una copia allí
    if output_dir:
        out_path = Path(output_dir)
        if out_path.exists() and out_path.is_dir():
            file_name = f"F004_GSW_DATO_{date.today().strftime('%Y-%m-%d')}.xlsx"
            (out_path / file_name).write_bytes(excel_bytes)

    return excel_bytes

